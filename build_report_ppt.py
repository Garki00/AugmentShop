from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_CONNECTOR
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "实践项目报告_跨境电商平台.pptx"
IMG = ROOT / "docs" / "screenshots"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(12, 18, 30)
PANEL = RGBColor(22, 32, 49)
PANEL2 = RGBColor(28, 42, 63)
WHITE = RGBColor(245, 248, 252)
MUTED = RGBColor(166, 181, 201)
CYAN = RGBColor(53, 211, 201)
BLUE = RGBColor(83, 140, 255)
ORANGE = RGBColor(255, 170, 75)
GREEN = RGBColor(91, 214, 137)
RED = RGBColor(246, 103, 103)

def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape

def text(slide, s, x, y, w, h, size=18, color=WHITE, bold=False, font="Aptos",
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box

def title(slide, kicker, heading, page):
    text(slide, kicker.upper(), 0.65, 0.35, 5.5, 0.25, 10, CYAN, True)
    text(slide, heading, 0.65, 0.65, 10.8, 0.55, 27, WHITE, True)
    text(slide, f"{page:02d}", 12.2, 0.42, 0.5, 0.3, 11, MUTED, True, align=PP_ALIGN.RIGHT)
    rect(slide, 0.65, 1.33, 12.0, 0.02, PANEL2)

def base():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    return slide

def bullet(slide, s, x, y, w, color=WHITE, size=16, accent=CYAN):
    rect(slide, x, y+0.09, 0.09, 0.09, accent, radius=True)
    text(slide, s, x+0.2, y, w-0.2, 0.34, size, color)

def metric(slide, value, label, x, y, w, accent=CYAN):
    rect(slide, x, y, w, 1.0, PANEL, radius=True)
    text(slide, value, x+0.18, y+0.12, w-0.36, 0.43, 25, accent, True)
    text(slide, label, x+0.18, y+0.61, w-0.36, 0.22, 11, MUTED)

def add_image(slide, filename, x, y, w, h):
    slide.shapes.add_picture(str(IMG / filename), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

# 1 cover
slide = base()
rect(slide, 0, 0, 13.333, 7.5, BG)
rect(slide, 8.7, 0, 4.633, 7.5, PANEL, line=PANEL)
for i, c in enumerate([CYAN, BLUE, ORANGE]):
    rect(slide, 9.25+i*0.85, 1.15+i*0.28, 0.56, 4.6-i*0.45, c, radius=True)
text(slide, "实践项目报告", 0.75, 1.0, 5.5, 0.4, 16, CYAN, True)
text(slide, "跨境电商平台\n与 AI 商品客服", 0.75, 1.65, 7.3, 1.55, 34, WHITE, True)
text(slide, "Flask + Vue 3 + MySQL + Docker Compose\n基于 RAG 的商品问答实践", 0.8, 3.55, 6.8, 0.75, 17, MUTED)
rect(slide, 0.8, 5.55, 2.15, 0.42, CYAN, radius=True)
text(slide, "项目实战汇报", 0.8, 5.63, 2.15, 0.2, 12, BG, True, align=PP_ALIGN.CENTER)
text(slide, "2026.09", 0.8, 6.45, 2, 0.25, 11, MUTED)

# 2 overview
slide = base(); title(slide, "01 / 项目概览", "把一个可运行的商城系统做完整", 2)
text(slide, "本项目面向跨境商品展示与咨询场景，完成了从数据种子、前后端接口到容器化部署和自动化验证的完整闭环。", 0.75, 1.65, 11.7, 0.55, 18, WHITE)
metric(slide, "100", "真实商品数据", 0.75, 2.65, 2.25, CYAN)
metric(slide, "11", "商品分类", 3.2, 2.65, 2.25, BLUE)
metric(slide, "37/37", "API 冒烟测试", 5.65, 2.65, 2.25, GREEN)
metric(slide, "24/24", "真实 UI 测试", 8.1, 2.65, 2.25, ORANGE)
metric(slide, "3", "Docker 服务", 10.55, 2.65, 2.05, CYAN)
bullet(slide, "用户侧：注册登录、商品搜索、分页浏览、详情查看、AI 咨询", 0.9, 4.35, 5.7)
bullet(slide, "管理侧：商品新增、编辑、删除、图片上传、Embedding 刷新", 0.9, 4.9, 5.7, accent=ORANGE)
bullet(slide, "工程侧：SQLite 本地开发，MySQL 生产容器，nginx 统一入口", 6.8, 4.35, 5.7, accent=BLUE)
bullet(slide, "验证侧：接口、浏览器 UI、构建、Compose 配置均有可复现实证", 6.8, 4.9, 5.7, accent=GREEN)

# 3 architecture
slide = base(); title(slide, "02 / 系统设计", "前后端分离，容器化交付", 3)
text(slide, "通过 /api 和 /static 约定隔离前端页面、业务 API 与静态资源，既便于本地开发，也支持 Docker 一键启动。", 0.75, 1.58, 11.7, 0.45, 16, MUTED)
rect(slide, 0.9, 2.55, 3.0, 2.55, PANEL, radius=True)
text(slide, "Vue 3 SPA", 1.15, 2.9, 2.5, 0.35, 20, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, "Vite · vue-router · axios\nHome / Detail / Admin\nProductCard · ChatPanel", 1.15, 3.55, 2.5, 0.9, 15, WHITE, align=PP_ALIGN.CENTER)
rect(slide, 5.15, 2.55, 3.0, 2.55, PANEL, radius=True)
text(slide, "Flask REST API", 5.38, 2.9, 2.55, 0.35, 20, BLUE, True, align=PP_ALIGN.CENTER)
text(slide, "SQLAlchemy · PyJWT\nAuth / Products / Admin / Chat\nGunicorn 4 workers", 5.38, 3.55, 2.55, 0.9, 15, WHITE, align=PP_ALIGN.CENTER)
rect(slide, 9.4, 2.55, 3.0, 2.55, PANEL, radius=True)
text(slide, "MySQL 8", 9.65, 2.9, 2.5, 0.35, 20, ORANGE, True, align=PP_ALIGN.CENTER)
text(slide, "users / products\nEmbedding LONGTEXT 缓存\n静态图片与上传文件", 9.65, 3.55, 2.5, 0.9, 15, WHITE, align=PP_ALIGN.CENTER)
for x in [3.95, 8.2]:
    rect(slide, x, 3.7, 0.8, 0.08, CYAN, radius=True)
text(slide, "nginx :8080 统一入口    /api → backend    /static → backend", 2.0, 5.85, 9.4, 0.35, 15, CYAN, True, align=PP_ALIGN.CENTER)

# 4 features screenshots
slide = base(); title(slide, "03 / 功能实现", "从商品浏览到管理闭环", 4)
add_image(slide, "01-home.png", 0.7, 1.7, 3.75, 2.25)
add_image(slide, "02-search.png", 4.8, 1.7, 3.75, 2.25)
add_image(slide, "05-admin.png", 8.9, 1.7, 3.75, 2.25)
text(slide, "首页推荐与分页", 0.7, 4.15, 3.75, 0.3, 16, WHITE, True, align=PP_ALIGN.CENTER)
text(slide, "关键词检索与结果恢复", 4.8, 4.15, 3.75, 0.3, 16, WHITE, True, align=PP_ALIGN.CENTER)
text(slide, "后台商品 CRUD 与统计", 8.9, 4.15, 3.75, 0.3, 16, WHITE, True, align=PP_ALIGN.CENTER)
bullet(slide, "商品列表默认 18 条/页，三列六行，共 6 页", 0.95, 5.1, 5.5)
bullet(slide, "推荐接口默认随机返回 5 个商品，前端 3.5 秒轮播", 0.95, 5.58, 5.5, accent=ORANGE)
bullet(slide, "后台接口通过 login_required + admin_required 实现 RBAC", 6.65, 5.1, 5.7, accent=BLUE)
bullet(slide, "图片支持本地静态资源和管理员上传，失败自动占位", 6.65, 5.58, 5.7, accent=GREEN)

# 5 auth/admin
slide = base(); title(slide, "04 / 安全与管理", "JWT 认证 + RBAC 权限控制", 5)
rect(slide, 0.8, 1.75, 5.7, 4.6, PANEL, radius=True)
text(slide, "认证流程", 1.15, 2.05, 2.0, 0.3, 20, CYAN, True)
steps = [("01", "注册 / 登录"), ("02", "Werkzeug 密码哈希"), ("03", "签发 HS256 JWT"), ("04", "Axios 自动携带 Bearer"), ("05", "接口按角色放行")]
for i, (n, s) in enumerate(steps):
    y = 2.65 + i*0.58
    rect(slide, 1.15, y, 0.42, 0.32, CYAN if i < 3 else BLUE, radius=True)
    text(slide, n, 1.15, y+0.06, 0.42, 0.15, 10, BG, True, align=PP_ALIGN.CENTER)
    text(slide, s, 1.8, y+0.02, 3.9, 0.25, 15, WHITE)
rect(slide, 6.85, 1.75, 5.7, 4.6, PANEL, radius=True)
text(slide, "RBAC 验证结果", 7.2, 2.05, 3.0, 0.3, 20, ORANGE, True)
metric(slide, "401", "未登录访问受保护接口", 7.2, 2.7, 2.25, RED)
metric(slide, "403", "普通用户访问管理接口", 9.7, 2.7, 2.25, ORANGE)
metric(slide, "201", "管理员成功创建商品", 7.2, 4.05, 2.25, GREEN)
metric(slide, "404", "删除后再次查询", 9.7, 4.05, 2.25, BLUE)
text(slide, "权限判断在后端完成，前端仅负责展示与路由体验。", 7.2, 5.55, 4.8, 0.4, 14, MUTED)

# 6 AI RAG
slide = base(); title(slide, "05 / AI 能力", "Embedding 检索 + 对话模型生成", 6)
add_image(slide, "04-detail-chat.png", 0.75, 1.65, 5.0, 3.0)
text(slide, "商品详情页嵌入 ChatPanel，回答中携带真实商品价格与属性，并返回相关商品。", 0.9, 4.85, 4.7, 0.65, 15, WHITE)
rect(slide, 6.35, 1.65, 6.2, 4.9, PANEL, radius=True)
text(slide, "一次问答的 5 个步骤", 6.75, 1.98, 3.5, 0.3, 20, CYAN, True)
rag = [
    ("1", "读取当前商品", "名称、价格、描述、分类、品牌"),
    ("2", "生成问题向量", "Qwen3-Embedding-8B"),
    ("3", "相似度检索", "从 100 个商品中取 Top 3"),
    ("4", "拼接受控 Prompt", "真实字段注入，限制价格幻觉"),
    ("5", "生成回答", "Qwen2.5-7B-Instruct + 降级兜底"),
]
for i, (n, h, d) in enumerate(rag):
    y = 2.55 + i*0.68
    rect(slide, 6.75, y, 0.38, 0.32, BLUE if i != 4 else ORANGE, radius=True)
    text(slide, n, 6.75, y+0.06, 0.38, 0.15, 10, BG, True, align=PP_ALIGN.CENTER)
    text(slide, h, 7.35, y, 2.0, 0.25, 15, WHITE, True)
    text(slide, d, 9.35, y, 2.75, 0.25, 13, MUTED)

# 7 deployment
slide = base(); title(slide, "06 / 部署与工程实践", "docker compose up -d --build 即可运行", 7)
text(slide, "部署链路围绕健康检查和启动顺序设计，解决了 MySQL 初始化、Embedding 预热和 nginx 502 等实际问题。", 0.75, 1.55, 11.8, 0.5, 16, MUTED)
for x, head, body, accent in [
    (0.9, "mysql", "MySQL 8\n3307 → 3306\nhealthcheck", ORANGE),
    (4.45, "backend", "Flask + Gunicorn\nprestart 初始化\n100 商品 + Embedding", BLUE),
    (8.0, "frontend", "Vite build\nnginx :80\n/api 与 /static 代理", CYAN),
]:
    rect(slide, x, 2.65, 3.0, 2.2, PANEL, radius=True)
    rect(slide, x, 2.65, 3.0, 0.12, accent)
    text(slide, head, x+0.22, 3.05, 2.55, 0.3, 20, accent, True)
    text(slide, body, x+0.22, 3.65, 2.55, 0.9, 15, WHITE)
for x in [3.95, 7.5]:
    rect(slide, x, 3.65, 0.42, 0.08, MUTED, radius=True)
text(slide, "关键工程处理", 0.9, 5.35, 2.6, 0.3, 18, WHITE, True)
bullet(slide, "backend 依赖 mysql healthy，避免启动阶段 502", 0.9, 5.8, 5.6, accent=GREEN)
bullet(slide, "预启动阶段集中生成 Embedding，避免 4 个 worker 重复请求", 6.7, 5.8, 5.8, accent=ORANGE)

# 8 testing
slide = base(); title(slide, "07 / 测试与结果", "以可复现实证证明项目可用", 8)
metric(slide, "37/37", "smoke_test API 用例", 0.85, 1.7, 2.45, GREEN)
metric(slide, "24/24", "ui_check 浏览器用例", 3.55, 1.7, 2.45, GREEN)
metric(slide, "104", "前端构建模块", 6.25, 1.7, 2.45, BLUE)
metric(slide, "100/100", "商品完成向量缓存", 8.95, 1.7, 2.45, CYAN)
add_image(slide, "06-admin-created.png", 0.9, 3.25, 3.45, 2.35)
add_image(slide, "07-admin-after-delete.png", 4.95, 3.25, 3.45, 2.35)
rect(slide, 9.0, 3.25, 3.3, 2.35, PANEL, radius=True)
text(slide, "验证覆盖", 9.3, 3.55, 2.6, 0.3, 18, CYAN, True)
bullet(slide, "分页与搜索", 9.3, 4.05, 2.5, size=13)
bullet(slide, "JWT / RBAC", 9.3, 4.48, 2.5, size=13, accent=BLUE)
bullet(slide, "CRUD / 上传", 9.3, 4.91, 2.5, size=13, accent=ORANGE)
bullet(slide, "AI 真实价格", 9.3, 5.34, 2.5, size=13, accent=GREEN)

# 9 issues
slide = base(); title(slide, "08 / 问题与改进", "实践中暴露的问题，也是工程价值所在", 9)
items = [
    ("图片布局", "图片尺寸与卡片容器耦合，导致商品图被裁切；改为固定 stage + object-fit: contain。", CYAN),
    ("分页竞态", "快速切换页码时请求返回顺序不确定；使用请求序号 reqSeq 丢弃过期响应。", BLUE),
    ("启动 502", "nginx 启动早于 backend；增加 healthcheck 与 depends_on 条件等待。", ORANGE),
    ("AI 成本", "Embedding 首次预热耗时较长；通过数据库缓存和预启动集中生成降低重复调用。", GREEN),
]
for i, (h, d, c) in enumerate(items):
    y = 1.75 + i*1.18
    rect(slide, 0.9, y, 11.6, 0.86, PANEL, radius=True)
    rect(slide, 0.9, y, 0.12, 0.86, c)
    text(slide, h, 1.25, y+0.17, 1.5, 0.25, 16, c, True)
    text(slide, d, 2.8, y+0.14, 9.0, 0.42, 15, WHITE)
text(slide, "这些修复被纳入 QA 文档与演示脚本，形成可复用的排障经验。", 0.95, 6.55, 11.2, 0.3, 15, MUTED, align=PP_ALIGN.CENTER)

# 10 summary
slide = base(); title(slide, "09 / 项目总结", "从功能实现走向可验证、可交付", 10)
text(slide, "本次实践的核心收获，不只是完成一个页面，而是建立了一套从需求到交付的工程闭环。", 0.85, 1.7, 11.5, 0.5, 19, WHITE)
for x, h, d, c in [
    (0.9, "完成度", "核心业务链路完整\n用户、管理员、AI 三类角色场景可演示", CYAN),
    (4.45, "可靠性", "API 与 UI 双层测试\n测试结果可复现、可追踪", GREEN),
    (8.0, "可演进", "后续可接入购物车、订单、i18n\n并升级向量数据库与限流", ORANGE),
]:
    rect(slide, x, 2.75, 3.0, 2.3, PANEL, radius=True)
    text(slide, h, x+0.25, 3.15, 2.5, 0.35, 21, c, True, align=PP_ALIGN.CENTER)
    text(slide, d, x+0.28, 3.9, 2.45, 0.85, 15, WHITE, align=PP_ALIGN.CENTER)
text(slide, "谢谢聆听", 0.9, 6.2, 11.5, 0.45, 24, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, "项目资料：README.md · docs/architecture.md · QA.md · tests/", 0.9, 6.75, 11.5, 0.25, 11, MUTED, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
